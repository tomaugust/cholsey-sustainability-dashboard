"""Tests for cholsey_pipeline.fetch.ons_population_dwellings.

Uses committed, real fixtures (a trimmed real ONS search-results page, a
trimmed real ad-hoc release landing page, a real-values parish-population
xlsx, and real nomis API CSV responses for 3 of Cholsey/Moulsford's real
OA21CD codes, P2.8) -- no network, per development-plan.md §5.1. The live
network calls (fetch_current_parish_population_url, fetch_parish_population,
fetch_oa_counts) are thin wrappers, same pattern as every other Phase 2
fetcher.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from cholsey_pipeline.fetch.ons_population_dwellings import (
    OnsDiscoveryError,
    ParishPopulationRecord,
    discover_parish_population_download_url,
    find_latest_parish_population_landing_path,
    parse_nomis_oa_csv,
    parse_parish_population_sheet,
)

FIXTURES_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "ons_population_dwellings"


class TestFindLatestParishPopulationLandingPath:
    def test_finds_newest_vintage_from_real_snippet(self) -> None:
        """Real ONS search results (2026-09-30): mid-2022, mid-2021, and a
        decoy Cornwall-specific release whose slug also contains
        "mid2022" but doesn't match the parish-estimates slug shape."""
        html = (FIXTURES_DIR / "search_results_sample.html").read_text()
        year, path = find_latest_parish_population_landing_path(html)
        assert year == 2022
        assert path.endswith(
            "2502parishpopulationestimatesformid2022basedonbestfittingofoutputareastoparishes"
        )

    def test_ignores_lookalike_cornwall_slug(self) -> None:
        """The Cornwall release's slug contains "mid2022" but in a
        different shape ("...forparishesincornwallmid2022") and must not
        be picked up as a real parish-population vintage."""
        html = (
            '<a href="/peoplepopulationandcommunity/populationandmigration/'
            "populationestimates/adhocs/"
            '2212onslocalpopulationestimatesbyfiveyearagebandforparishesincornwallmid2022">x</a>'
        )
        with pytest.raises(OnsDiscoveryError):
            find_latest_parish_population_landing_path(html)

    def test_no_match_raises(self) -> None:
        with pytest.raises(OnsDiscoveryError, match="No 'parish population"):
            find_latest_parish_population_landing_path("<html>nothing</html>")


class TestDiscoverParishPopulationDownloadUrl:
    def test_finds_real_xlsx_link(self) -> None:
        html = (FIXTURES_DIR / "landing_page_sample.html").read_text()
        url = discover_parish_population_download_url(html)
        assert url == (
            "https://www.ons.gov.uk/file?uri=/peoplepopulationandcommunity/"
            "populationandmigration/populationestimates/adhocs/"
            "2502parishpopulationestimatesformid2022basedonbestfittingofoutputareastoparishes/"
            "parishmid2022popest.xlsx"
        )

    def test_no_match_raises(self) -> None:
        with pytest.raises(OnsDiscoveryError):
            discover_parish_population_download_url("<html></html>")


class TestParseParishPopulationSheet:
    def test_real_cholsey_and_moulsford_values(self) -> None:
        """Real mid-2022 values verified live 2026-09-30 while building
        this module: Cholsey 4,423 (vs mid-2021's 4,404), Moulsford 619."""
        workbook = load_workbook(
            FIXTURES_DIR / "parish_population_mid2022_sample.xlsx", data_only=True
        )
        worksheet = workbook["Parish Populations"]
        rows = list(worksheet.iter_rows(values_only=True))
        records = parse_parish_population_sheet(rows, 2022)
        by_code = {r.parish_code: r for r in records}
        assert by_code["E04012474"].total_population == 4423
        assert by_code["E04012474"].parish_name == "Cholsey"
        assert by_code["E04008148"].total_population == 619
        assert all(r.vintage_year == 2022 for r in records)

    def test_filters_to_requested_codes(self) -> None:
        workbook = load_workbook(
            FIXTURES_DIR / "parish_population_mid2022_sample.xlsx", data_only=True
        )
        rows = list(workbook["Parish Populations"].iter_rows(values_only=True))
        records = parse_parish_population_sheet(rows, 2022, parish_codes={"E04012474"})
        assert [r.parish_code for r in records] == ["E04012474"]

    def test_total_population_is_int(self) -> None:
        workbook = load_workbook(
            FIXTURES_DIR / "parish_population_mid2022_sample.xlsx", data_only=True
        )
        rows = list(workbook["Parish Populations"].iter_rows(values_only=True))
        records = parse_parish_population_sheet(rows, 2022)
        assert all(isinstance(r.total_population, int) for r in records)

    def test_future_vintage_column_names_are_matched_by_pattern(self) -> None:
        """Regression test for the fixed bug: a hardcoded "PAR22CD" column
        name would raise on a later vintage keyed to a newer parish
        edition (e.g. mid-2023's PAR23CD) -- the pattern match must find
        it regardless of the specific year suffix."""
        rows = [
            ("PAR23CD", "PAR23NM", "Total"),
            ("E04012474", "Cholsey", 4500),
        ]
        records = parse_parish_population_sheet(rows, 2023)
        assert records == [
            ParishPopulationRecord(
                vintage_year=2023,
                parish_code="E04012474",
                parish_name="Cholsey",
                total_population=4500,
            )
        ]

    def test_blank_footnote_row_is_skipped(self) -> None:
        """A trailing blank/footnote row (no parish code) must not become
        a record with parish_code=None."""
        rows = [
            ("PAR22CD", "PAR22NM", "Total"),
            ("E04012474", "Cholsey", 4423),
            (None, "Source: ONS", None),
        ]
        records = parse_parish_population_sheet(rows, 2022)
        assert len(records) == 1
        assert records[0].parish_code == "E04012474"


class TestParseNomisOaCsv:
    def test_real_population_totals_exclude_breakdown_rows(self) -> None:
        """Real live-fetched values (2026-09-30) for 3 of Cholsey/
        Moulsford's real OA21CD codes -- only the "Total: All usual
        residents" row per OA should be kept, not the "lives in a
        household"/"lives in a communal establishment" breakdown rows."""
        csv_text = (FIXTURES_DIR / "nomis_population_sample.csv").read_text()
        counts = parse_nomis_oa_csv(csv_text, "population")
        assert counts == {
            "E00145768": 358,
            "E00145777": 283,
            "E00186067": 254,
        }

    def test_real_household_totals(self) -> None:
        csv_text = (FIXTURES_DIR / "nomis_households_sample.csv").read_text()
        counts = parse_nomis_oa_csv(csv_text, "households")
        assert counts == {
            "E00145768": 135,
            "E00145777": 113,
            "E00186067": 95,
        }
