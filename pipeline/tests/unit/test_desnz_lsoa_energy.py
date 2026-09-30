"""Tests for cholsey_pipeline.fetch.desnz_lsoa_energy.

Uses committed, real fixtures (trimmed 2023 sheets from the live DESNZ
LSOA-level and regional/local-authority-level electricity releases,
keeping Cholsey's 3 LSOAs and South Oxfordshire/England's regional rows,
exactly as fetched while building this module, P2.5) -- no network, per
development-plan.md §5.1. The live network calls (fetch_current_download_url,
fetch_lsoa_energy, fetch_regional_la_energy) are thin wrappers, same
pattern as geography.boundaries.fetch_boundary.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from cholsey_pipeline.fetch.desnz_lsoa_energy import (
    LSOA_FILENAME_PATTERNS,
    REGIONAL_LA_FILENAME_PATTERNS,
    DesnzDiscoveryError,
    discover_download_url,
    parse_lsoa_sheet,
    parse_regional_la_sheet,
    read_lsoa_energy,
    read_regional_la_energy,
)

FIXTURES_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "desnz_lsoa_energy"
CHOLSEY_LSOAS = {"E01035751", "E01028619", "E01035752"}


def _rows_from(path: Path, sheet: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    return list(wb[sheet].iter_rows(values_only=True))


class TestDiscoverDownloadUrl:
    def test_finds_matching_link(self) -> None:
        html = """
        <a href="https://assets.publishing.service.gov.uk/media/abc123/MSOA_domestic_elec_2010-2024.xlsx">MSOA</a>
        <a href="https://assets.publishing.service.gov.uk/media/def456/LSOA_domestic_elec_2010-2024.xlsx">LSOA</a>
        """
        url = discover_download_url(html, LSOA_FILENAME_PATTERNS["electricity"])
        assert url == (
            "https://assets.publishing.service.gov.uk/media/def456/"
            "LSOA_domestic_elec_2010-2024.xlsx"
        )

    def test_does_not_match_gas_pattern_against_electricity_link(self) -> None:
        html = (
            '<a href="https://assets.publishing.service.gov.uk/media/def456/'
            'LSOA_domestic_elec_2010-2024.xlsx">LSOA</a>'
        )
        with pytest.raises(DesnzDiscoveryError, match="No .xlsx link"):
            discover_download_url(html, LSOA_FILENAME_PATTERNS["gas"])

    def test_no_match_raises(self) -> None:
        with pytest.raises(DesnzDiscoveryError):
            discover_download_url(
                "<html>nothing here</html>", LSOA_FILENAME_PATTERNS["electricity"]
            )

    def test_regional_la_pattern_matches_real_filename(self) -> None:
        html = (
            '<a href="https://assets.publishing.service.gov.uk/media/xyz/'
            'Subnational_electricity_consumption_statistics_2005-2024.xlsx">x</a>'
        )
        url = discover_download_url(html, REGIONAL_LA_FILENAME_PATTERNS["electricity"])
        assert url.endswith("Subnational_electricity_consumption_statistics_2005-2024.xlsx")


class TestParseLsoaSheet:
    def test_real_cholsey_lsoa_values(self) -> None:
        """Real values verified live 2026-09-29 while building this
        module."""
        rows = _rows_from(FIXTURES_DIR / "lsoa_electricity_2023_sample.xlsx", "2023")
        records = {r.lsoa_code: r for r in parse_lsoa_sheet(rows, 2023, "electricity")}
        assert records["E01028619"].number_of_meters == 795
        assert records["E01028619"].la_code == "E07000179"
        assert records["E01028619"].la_name == "South Oxfordshire"
        assert records["E01035751"].total_consumption_kwh == pytest.approx(2598587.06, abs=0.1)
        assert records["E01035752"].lsoa_name == "South Oxfordshire 015I"
        assert records["E01028619"].fuel == "electricity"

    def test_filters_to_requested_lsoa_codes(self) -> None:
        rows = _rows_from(FIXTURES_DIR / "lsoa_electricity_2023_sample.xlsx", "2023")
        records = parse_lsoa_sheet(rows, 2023, "electricity", lsoa_codes={"E01028619"})
        assert {r.lsoa_code for r in records} == {"E01028619"}

    def test_year_is_attached(self) -> None:
        rows = _rows_from(FIXTURES_DIR / "lsoa_electricity_2023_sample.xlsx", "2023")
        records = parse_lsoa_sheet(rows, 2023, "electricity")
        assert all(r.year == 2023 for r in records)

    def test_bad_header_raises(self) -> None:
        from cholsey_pipeline.fetch.desnz_lsoa_energy import DesnzParseError

        rows = _rows_from(FIXTURES_DIR / "lsoa_electricity_2023_sample.xlsx", "2023")
        mutated = list(rows)
        header = list(mutated[4])
        header[4] = "Some other column"  # was "LSOA code"
        mutated[4] = tuple(header)
        with pytest.raises(DesnzParseError, match="lsoa code"):
            parse_lsoa_sheet(mutated, 2023, "electricity")


class TestParseRegionalLaSheet:
    def test_real_south_oxfordshire_and_england_values(self) -> None:
        """Real values verified live 2026-09-29."""
        rows = _rows_from(FIXTURES_DIR / "regional_la_electricity_2023_sample.xlsx", "2023")
        records = {r.area_code: r for r in parse_regional_la_sheet(rows, 2023, "electricity")}
        south_oxon = records["E07000179"]
        assert south_oxon.area_name == "South Oxfordshire"
        assert south_oxon.total_domestic_consumption_gwh == pytest.approx(272.26, abs=0.01)
        england = records["E92000001"]
        assert england.area_name == "England"
        assert england.total_domestic_consumption_gwh == pytest.approx(83110.58, abs=0.1)
        assert south_oxon.fuel == "electricity"

    def test_filters_to_requested_area_codes(self) -> None:
        rows = _rows_from(FIXTURES_DIR / "regional_la_electricity_2023_sample.xlsx", "2023")
        records = parse_regional_la_sheet(rows, 2023, "electricity", area_codes={"E92000001"})
        assert {r.area_code for r in records} == {"E92000001"}


class TestReadLsoaEnergy:
    def test_reads_named_sheet(self) -> None:
        records = read_lsoa_energy(
            str(FIXTURES_DIR / "lsoa_electricity_2023_sample.xlsx"),
            [2023],
            "electricity",
            CHOLSEY_LSOAS,
        )
        assert len(records) == 3

    def test_missing_sheet_raises(self) -> None:
        from cholsey_pipeline.fetch.desnz_lsoa_energy import DesnzParseError

        with pytest.raises(DesnzParseError, match="2022"):
            read_lsoa_energy(
                str(FIXTURES_DIR / "lsoa_electricity_2023_sample.xlsx"), [2022], "electricity"
            )


class TestReadRegionalLaEnergy:
    def test_reads_named_sheet(self) -> None:
        records = read_regional_la_energy(
            str(FIXTURES_DIR / "regional_la_electricity_2023_sample.xlsx"),
            [2023],
            "electricity",
            {"E07000179", "E92000001"},
        )
        assert len(records) == 2
