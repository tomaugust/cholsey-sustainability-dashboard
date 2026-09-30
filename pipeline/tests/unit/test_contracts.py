"""Tests for cholsey_pipeline.contracts (P2.10).

Tests of success (development-plan.md §3 Phase 2):
  - each source's declared contract validates its own real fetcher output
    against a small real fixture (via the source's own parser, reusing
    each fetcher's existing committed fixture rather than duplicating it).
  - negative: a renamed column, a dropped column, and a -50% row count
    each cause a contract failure naming the source.
  - idempotence: running validate_schema twice against the same records
    doesn't mutate anything or behave differently the second time.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cholsey_pipeline.contracts import (
    SOURCE_CONTRACTS,
    ContractViolation,
    SchemaContract,
    validate_row_count,
    validate_schema,
)
from cholsey_pipeline.fetch.desnz_lsoa_energy import parse_lsoa_sheet
from cholsey_pipeline.fetch.desnz_postcode_energy import parse_postcode_csv
from cholsey_pipeline.fetch.dluhc_epc_register import parse_domestic_search_response
from cholsey_pipeline.fetch.forest_research_canopy import parse_canopy_response
from cholsey_pipeline.fetch.ons_population_dwellings import parse_parish_population_sheet

FIXTURES_ROOT = Path(__file__).resolve().parents[1] / "fixtures"


@dataclass(frozen=True)
class _FakeLsoaRecord:
    """A record shaped like LsoaEnergyRecord but with a renamed field, to
    test that validate_schema catches a renamed/dropped column without
    needing a real broken upstream file."""

    year: int
    lsoa_code: str
    lsoa_name: str
    la_code: str
    la_name: str
    meters_count: float  # renamed from "number_of_meters"
    total_consumption_kwh: float


class TestValidateSchemaAgainstRealFetcherOutput:
    def test_desnz_lsoa_energy_real_fixture_passes(self) -> None:
        fixture_dir = FIXTURES_ROOT / "desnz_lsoa_energy"
        wb = load_workbook(fixture_dir / "lsoa_electricity_2023_sample.xlsx", data_only=True)
        rows = list(wb["2023"].iter_rows(values_only=True))
        records = parse_lsoa_sheet(rows, 2023)
        validate_schema(records, SOURCE_CONTRACTS["desnz_lsoa_energy"])

    def test_desnz_postcode_energy_real_fixture_passes(self) -> None:
        fixture_dir = FIXTURES_ROOT / "desnz_postcode_energy"
        csv_text = (fixture_dir / "ox10_electricity_2024_sample.csv").read_text()
        records = parse_postcode_csv(csv_text, 2024)
        validate_schema(records, SOURCE_CONTRACTS["desnz_postcode_energy"])

    def test_ons_parish_population_real_fixture_passes(self) -> None:
        fixture_dir = FIXTURES_ROOT / "ons_population_dwellings"
        wb = load_workbook(fixture_dir / "parish_population_mid2022_sample.xlsx", data_only=True)
        rows = list(wb["Parish Populations"].iter_rows(values_only=True))
        records = parse_parish_population_sheet(rows, 2022)
        validate_schema(records, SOURCE_CONTRACTS["ons_parish_population"])

    def test_dluhc_epc_register_fixture_passes(self) -> None:
        fixture_dir = FIXTURES_ROOT / "dluhc_epc_register"
        payload = json.loads((fixture_dir / "domestic_search_response_sample.json").read_text())
        records = parse_domestic_search_response(payload)
        validate_schema(records, SOURCE_CONTRACTS["dluhc_epc_register"])

    def test_forest_research_canopy_real_fixture_passes(self) -> None:
        fixture_dir = FIXTURES_ROOT / "forest_research_canopy"
        payload = json.loads((fixture_dir / "cholsey_wallingford_sample.json").read_text())
        records = parse_canopy_response(payload)
        validate_schema(records, SOURCE_CONTRACTS["forest_research_canopy"])


class TestValidateSchemaNegative:
    def test_renamed_column_fails(self) -> None:
        """A renamed column means the contract's expected field name is
        simply absent from the record type -- caught the same way as a
        dropped column."""
        fake_records = [
            _FakeLsoaRecord(
                2023, "E01028619", "Cholsey 1", "E07000179", "South Oxfordshire", 795, 2852777.0
            )
        ]
        with pytest.raises(ContractViolation, match="number_of_meters"):
            validate_schema(fake_records, SOURCE_CONTRACTS["desnz_lsoa_energy"])

    def test_null_key_field_fails(self) -> None:
        contract = SchemaContract(
            source_id="test_source", expected_fields=("code", "value"), key_fields=("code",)
        )

        @dataclass(frozen=True)
        class Row:
            code: str | None
            value: int

        with pytest.raises(ContractViolation, match="null"):
            validate_schema([Row(None, 1)], contract)

    def test_duplicate_key_fails(self) -> None:
        contract = SchemaContract(
            source_id="test_source", expected_fields=("code", "value"), key_fields=("code",)
        )

        @dataclass(frozen=True)
        class Row:
            code: str
            value: int

        with pytest.raises(ContractViolation, match="duplicate"):
            validate_schema([Row("A", 1), Row("A", 2)], contract)

    def test_empty_batch_does_not_raise(self) -> None:
        """An empty batch isn't itself a schema violation -- that's
        validate_row_count's job (a batch of zero vs a previous non-zero
        count would exceed any reasonable tolerance)."""
        validate_schema([], SOURCE_CONTRACTS["desnz_lsoa_energy"])


class TestValidateRowCount:
    def test_no_previous_run_is_not_checked(self) -> None:
        contract = SOURCE_CONTRACTS["desnz_lsoa_energy"]
        validate_row_count(contract, current_count=3, previous_count=None)
        validate_row_count(contract, current_count=3, previous_count=0)

    def test_within_tolerance_passes(self) -> None:
        contract = SOURCE_CONTRACTS["desnz_lsoa_energy"]  # 30% tolerance
        validate_row_count(contract, current_count=90, previous_count=100)

    def test_minus_50_percent_fails(self) -> None:
        """development-plan.md's own negative test: a -50% row count
        change must cause a hard failure naming the source."""
        contract = SOURCE_CONTRACTS["desnz_lsoa_energy"]  # 30% tolerance
        with pytest.raises(ContractViolation, match="desnz_lsoa_energy"):
            validate_row_count(contract, current_count=50, previous_count=100)

    def test_large_increase_also_fails(self) -> None:
        contract = SOURCE_CONTRACTS["desnz_lsoa_energy"]
        with pytest.raises(ContractViolation):
            validate_row_count(contract, current_count=200, previous_count=100)


class TestIdempotence:
    def test_validate_schema_twice_is_identical(self) -> None:
        fixture_dir = FIXTURES_ROOT / "desnz_postcode_energy"
        csv_text = (fixture_dir / "ox10_electricity_2024_sample.csv").read_text()
        records = parse_postcode_csv(csv_text, 2024)
        contract = SOURCE_CONTRACTS["desnz_postcode_energy"]
        validate_schema(records, contract)
        validate_schema(records, contract)  # no exception, no state carried over
