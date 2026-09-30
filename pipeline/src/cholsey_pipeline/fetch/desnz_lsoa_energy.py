"""Fetch DESNZ sub-national domestic electricity/gas consumption
(development-plan.md Phase 2, P2.5): LSOA-level figures for every
available year, plus district and national totals.

Real finding while building this module: the LSOA-level release
("Lower and middle super output areas <fuel> consumption") and the
district/national-level release ("Regional and local authority <fuel>
consumption statistics") are two DIFFERENT GOV.UK statistics
publications, not one release as P2.5's summary wording suggested --
related, in the same "sub-national electricity consumption data"
collection, but published and updated separately, each with its own
landing page and its own current-year `.xlsx` asset link. This module
treats them as the two sibling sources they actually are.

Both releases' asset URLs are hash-named
(assets.publishing.service.gov.uk/media/<hash>/...) and change on every
annual update, so this module discovers the current link by scraping the
landing page each run rather than hardcoding a snapshot (see
config/sources.yaml's discovery_rule note and P1's boundaries.py /
lsoa_overlap.py worklog entries for the same "search, don't guess"
pattern).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any, Literal

import requests
from openpyxl import load_workbook

from cholsey_pipeline.contracts import SOURCE_CONTRACTS, validate_row_count, validate_schema
from cholsey_pipeline.fetch.http import fetch_file, previous_row_count, record_row_count

Fuel = Literal["electricity", "gas"]

LSOA_LANDING_URLS: dict[Fuel, str] = {
    "electricity": (
        "https://www.gov.uk/government/statistics/"
        "lower-and-middle-super-output-areas-electricity-consumption"
    ),
    "gas": (
        "https://www.gov.uk/government/statistics/"
        "lower-and-middle-super-output-areas-gas-consumption"
    ),
}
REGIONAL_LA_LANDING_URLS: dict[Fuel, str] = {
    "electricity": (
        "https://www.gov.uk/government/statistics/"
        "regional-and-local-authority-electricity-consumption-statistics"
    ),
    "gas": (
        "https://www.gov.uk/government/statistics/"
        "regional-and-local-authority-gas-consumption-statistics"
    ),
}
"""Verified live 2026-09-29 (P2.5). Discovered via the GOV.UK site search
API (api/search.json), not guessed -- see the P2.5 worklog entry."""

LSOA_FILENAME_PATTERNS: dict[Fuel, re.Pattern[str]] = {
    "electricity": re.compile(r"LSOA_domestic_elec_\d{4}-\d{4}\.xlsx$"),
    "gas": re.compile(r"LSOA_domestic_gas_\d{4}-\d{4}\.xlsx$"),
}
REGIONAL_LA_FILENAME_PATTERNS: dict[Fuel, re.Pattern[str]] = {
    "electricity": re.compile(r"Subnational_electricity_consumption_statistics_\d{4}-\d{4}\.xlsx$"),
    "gas": re.compile(r"Subnational_gas_consumption_statistics_\d{4}-\d{4}\.xlsx$"),
}

ASSET_LINK_PATTERN = re.compile(
    r'href="(https://assets\.publishing\.service\.gov\.uk/media/[^"]+\.xlsx)"'
)

LSOA_HEADER_ROW = 5
"""1-indexed row the column headers are on, in every LSOA-level sheet --
verified live 2026-09-29 against the real 2023 electricity sheet."""

REGIONAL_LA_HEADER_ROW: dict[Fuel, int] = {"electricity": 5, "gas": 6}
"""1-indexed header row for the regional/local-authority sheets --
per-fuel, NOT the same for both. Real finding (2026-09-30, live-verified
against the real 2024 workbooks while building P3.4): gas's sheet has an
extra explanatory note row (about its extra "Notes" column, which
electricity's sheet doesn't have) above the header, pushing gas's header
down to row 6 vs electricity's row 5. P2.5 only ever verified the
electricity regional/LA sheet live -- this gas-specific layout was never
actually exercised until P3.4 tried to fetch it for real."""


class DesnzDiscoveryError(RuntimeError):
    """Raised when a GOV.UK landing page doesn't contain a matching .xlsx
    link -- never guessed, so a page layout change surfaces loudly."""


class DesnzParseError(RuntimeError):
    """Raised when a workbook's sheet doesn't have the expected columns."""


def _check_header(header_row: tuple[Any, ...], expected: dict[int, str], sheet_label: str) -> None:
    """Verify specific column positions' header text before reading data
    by position. Both parsers below read columns by fixed index (the real
    column order verified live 2026-09-29), so a column DESNZ inserts,
    removes or reorders upstream would otherwise silently shift values
    into the wrong fields while parsing still "succeeds" -- this makes
    that fail loudly instead, naming the sheet and the mismatched column.

    Matching is normalised (lowercased, whitespace-collapsed) and by
    substring, not exact equality, since these headers wrap onto multiple
    lines in the real workbook (e.g. "Total \\nconsumption\\n(kWh)").
    """
    for index, expected_substring in expected.items():
        actual = header_row[index] if index < len(header_row) else None
        normalised = " ".join(str(actual or "").lower().split())
        if expected_substring not in normalised:
            raise DesnzParseError(
                f"{sheet_label}: expected column {index} to contain "
                f"{expected_substring!r}, found {actual!r} -- the upstream "
                "column layout may have changed"
            )


def discover_download_url(landing_page_html: str, filename_pattern: re.Pattern[str]) -> str:
    """Find the current `.xlsx` asset URL on a GOV.UK statistics landing
    page's HTML, matching `filename_pattern`. Pure function -- tested
    against a committed fixture HTML snippet, not fetched live
    (development-plan.md §5.1)."""
    for match in ASSET_LINK_PATTERN.finditer(landing_page_html):
        url = match.group(1)
        if filename_pattern.search(url):
            return url
    raise DesnzDiscoveryError(
        f"No .xlsx link matching {filename_pattern.pattern!r} found on the landing page"
    )


def fetch_current_download_url(landing_url: str, filename_pattern: re.Pattern[str]) -> str:
    """Live version of discover_download_url: fetches the landing page
    then discovers the link. The network call itself isn't covered by
    tests, same as geography.boundaries.fetch_boundary."""
    response = requests.get(landing_url, timeout=30)
    response.raise_for_status()
    return discover_download_url(response.text, filename_pattern)


@dataclass(frozen=True)
class LsoaEnergyRecord:
    """One LSOA's domestic consumption for one year and fuel."""

    fuel: Fuel
    year: int
    lsoa_code: str
    lsoa_name: str
    la_code: str
    la_name: str
    number_of_meters: float
    total_consumption_kwh: float


@dataclass(frozen=True)
class AreaEnergyRecord:
    """One region/country/local-authority's domestic consumption for one
    year and fuel, from the regional/LA-level release."""

    fuel: Fuel
    year: int
    area_code: str
    area_name: str
    number_of_domestic_meters_thousands: float
    total_domestic_consumption_gwh: float


def _iter_data_rows(rows: list[tuple[Any, ...]], header_row: int) -> list[tuple[Any, ...]]:
    """Rows are 0-indexed in `rows`; `header_row` is 1-indexed (matching
    Excel's own row numbering, as used in this module's *_HEADER_ROW
    constants) -- data starts immediately after it."""
    return rows[header_row:]


_LSOA_EXPECTED_HEADER = {
    0: "local authority code",
    4: "lsoa code",
    5: "lower layer super output area",
    6: "number",
    7: "total",
}
_REGIONAL_LA_EXPECTED_HEADER: dict[Fuel, dict[int, str]] = {
    "electricity": {
        0: "code",
        1: "country or region",
        2: "local authority",
        5: "number of meters",
        10: "total consumption",
    },
    "gas": {
        0: "code",
        1: "country or region",
        2: "local authority",
        3: "notes",
        4: "number of meters",
        7: "total consumption",
    },
}
_REGIONAL_LA_DOMESTIC_METERS_COL: dict[Fuel, int] = {"electricity": 5, "gas": 4}
_REGIONAL_LA_DOMESTIC_CONSUMPTION_COL: dict[Fuel, int] = {"electricity": 10, "gas": 7}
"""Which column holds "All Domestic" meters (thousands) / "Domestic"
consumption (GWh) in the regional/LA sheet -- per-fuel, since gas's sheet
has an extra "Notes" column (electricity doesn't) and doesn't split
domestic meters into Standard/E7 sub-columns the way electricity does, so
its domestic totals land 1-3 columns earlier than electricity's. Real
finding, 2026-09-30 (P3.4) -- see REGIONAL_LA_HEADER_ROW's docstring."""


def parse_lsoa_sheet(
    rows: list[tuple[Any, ...]], year: int, fuel: Fuel, lsoa_codes: set[str] | None = None
) -> list[LsoaEnergyRecord]:
    """Parse one year's LSOA-level sheet (as a list of row tuples, exactly
    as `worksheet.iter_rows(values_only=True)` yields them) into
    LsoaEnergyRecords, optionally filtered to `lsoa_codes`.

    Column order verified live 2026-09-29 against the real 2023
    electricity sheet: la_code, la_name, msoa_code, msoa_name, lsoa_code,
    lsoa_name, num_meters, total_kwh, mean_kwh, median_kwh. The header row
    is checked against this before reading any data row by position, so a
    column DESNZ inserts/reorders upstream fails loudly instead of
    silently shifting values into the wrong field.
    """
    header = rows[LSOA_HEADER_ROW - 1]
    _check_header(header, _LSOA_EXPECTED_HEADER, f"LSOA {fuel} sheet")
    data_rows = _iter_data_rows(rows, LSOA_HEADER_ROW)
    records = []
    for row in data_rows:
        if len(row) < 8 or row[4] is None:
            continue
        lsoa_code = row[4]
        if lsoa_codes is not None and lsoa_code not in lsoa_codes:
            continue
        records.append(
            LsoaEnergyRecord(
                fuel=fuel,
                year=year,
                lsoa_code=lsoa_code,
                lsoa_name=row[5],
                la_code=row[0],
                la_name=row[1],
                number_of_meters=row[6],
                total_consumption_kwh=row[7],
            )
        )
    return records


def parse_regional_la_sheet(
    rows: list[tuple[Any, ...]], year: int, fuel: Fuel, area_codes: set[str] | None = None
) -> list[AreaEnergyRecord]:
    """Parse one year's regional/local-authority sheet into
    AreaEnergyRecords, optionally filtered to `area_codes` (GSS codes,
    e.g. a local authority district code or a country code like
    E92000001 for England).

    Column order verified live against the real 2024 sheets, PER FUEL --
    electricity (2026-09-29) and gas (2026-09-30) have genuinely
    different layouts (gas has an extra "Notes" column and doesn't split
    domestic meters into Standard/E7 sub-columns), so both the header row
    and the domestic meters/consumption column indices are looked up by
    `fuel`, not shared. Column 0 is always the area code, 1/2 the name
    (region name if this is a country/region summary row, else the local
    authority name). The header row is checked against the fuel's
    expected layout before reading any data row by position.
    """
    header_row = REGIONAL_LA_HEADER_ROW[fuel]
    header = rows[header_row - 1]
    _check_header(header, _REGIONAL_LA_EXPECTED_HEADER[fuel], f"regional/LA {fuel} sheet")
    data_rows = _iter_data_rows(rows, header_row)
    meters_col = _REGIONAL_LA_DOMESTIC_METERS_COL[fuel]
    consumption_col = _REGIONAL_LA_DOMESTIC_CONSUMPTION_COL[fuel]
    records = []
    for row in data_rows:
        if len(row) <= consumption_col or row[0] is None:
            continue
        area_code = row[0]
        if area_codes is not None and area_code not in area_codes:
            continue
        area_name = row[2] if row[2] and row[2] != "All local authorities" else row[1]
        records.append(
            AreaEnergyRecord(
                fuel=fuel,
                year=year,
                area_code=area_code,
                area_name=area_name,
                number_of_domestic_meters_thousands=row[meters_col],
                total_domestic_consumption_gwh=row[consumption_col],
            )
        )
    return records


def read_lsoa_energy(
    xlsx_path: str,
    years: list[int],
    fuel: Fuel,
    lsoa_codes: set[str] | None = None,
) -> list[LsoaEnergyRecord]:
    """Read and parse multiple years' sheets from a downloaded LSOA-level
    workbook. Each sheet is named after its year (e.g. "2023")."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    all_records = []
    for year in years:
        sheet_name = str(year)
        if sheet_name not in workbook.sheetnames:
            raise DesnzParseError(f"No sheet named {sheet_name!r} in {xlsx_path}")
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        all_records.extend(parse_lsoa_sheet(rows, year, fuel, lsoa_codes))
    return all_records


def read_regional_la_energy(
    xlsx_path: str,
    years: list[int],
    fuel: Fuel,
    area_codes: set[str] | None = None,
) -> list[AreaEnergyRecord]:
    """Read and parse multiple years' sheets from a downloaded regional/
    local-authority workbook."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    all_records = []
    for year in years:
        sheet_name = str(year)
        if sheet_name not in workbook.sheetnames:
            raise DesnzParseError(f"No sheet named {sheet_name!r} in {xlsx_path}")
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        all_records.extend(parse_regional_la_sheet(rows, year, fuel, area_codes))
    return all_records


def fetch_lsoa_energy(
    fuel: Fuel,
    years: list[int],
    lsoa_codes: set[str] | None = None,
) -> list[LsoaEnergyRecord]:
    """Full live fetch: discover the current LSOA-level download URL,
    fetch it (via fetch.http.fetch_file), and parse the requested years.
    Validates the result against `contracts.SOURCE_CONTRACTS` (schema and
    row-count-vs-previous-run) before returning -- development-plan.md
    P2.10's "refuses to proceed when a source changes shape" guarantee,
    enforced here rather than only in tests.

    The manifest `source_id` is `desnz_lsoa_energy/<fuel>` -- nested under
    the registry id (`config/sources.yaml`'s `desnz_lsoa_energy` entry) so
    provenance/licence lookups can find it, with the fuel as a distinct
    manifest history per fuel (electricity and gas are fetched, and their
    row counts compared, independently).
    """
    manifest_source_id = f"desnz_lsoa_energy/{fuel}"
    # A row-count baseline is only valid for the SAME years/lsoa_codes --
    # fetching one year then a decade of years is a different query, not a
    # real change in the data (cycle-2 review finding).
    query_signature = (
        f"years={sorted(years)}:lsoa_codes={sorted(lsoa_codes) if lsoa_codes else 'all'}"
    )
    download_url = fetch_current_download_url(LSOA_LANDING_URLS[fuel], LSOA_FILENAME_PATTERNS[fuel])
    previous_count = previous_row_count(manifest_source_id, query_signature)
    result = fetch_file(manifest_source_id, download_url, dest_filename=f"lsoa_{fuel}.xlsx")
    if result.file_path is None:
        raise RuntimeError(f"fetch_file returned no file_path for {manifest_source_id}")
    records = read_lsoa_energy(str(result.file_path), years, fuel, lsoa_codes)
    contract = SOURCE_CONTRACTS["desnz_lsoa_energy"]
    validate_schema(records, contract)
    validate_row_count(contract, len(records), previous_count)
    record_row_count(manifest_source_id, len(records), query_signature)
    return records


def fetch_regional_la_energy(
    fuel: Fuel,
    years: list[int],
    area_codes: set[str] | None = None,
) -> list[AreaEnergyRecord]:
    """Full live fetch: discover the current regional/LA-level download
    URL, fetch it, and parse the requested years. Same contract-validation
    and manifest-nesting pattern as `fetch_lsoa_energy`."""
    manifest_source_id = f"desnz_regional_la_energy/{fuel}"
    query_signature = (
        f"years={sorted(years)}:area_codes={sorted(area_codes) if area_codes else 'all'}"
    )
    download_url = fetch_current_download_url(
        REGIONAL_LA_LANDING_URLS[fuel], REGIONAL_LA_FILENAME_PATTERNS[fuel]
    )
    previous_count = previous_row_count(manifest_source_id, query_signature)
    result = fetch_file(manifest_source_id, download_url, dest_filename=f"regional_la_{fuel}.xlsx")
    if result.file_path is None:
        raise RuntimeError(f"fetch_file returned no file_path for {manifest_source_id}")
    records = read_regional_la_energy(str(result.file_path), years, fuel, area_codes)
    contract = SOURCE_CONTRACTS["desnz_regional_la_energy"]
    validate_schema(records, contract)
    validate_row_count(contract, len(records), previous_count)
    record_row_count(manifest_source_id, len(records), query_signature)
    return records
