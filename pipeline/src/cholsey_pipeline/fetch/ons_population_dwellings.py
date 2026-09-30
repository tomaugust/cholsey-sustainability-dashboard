"""Fetch ONS population and household/dwelling denominators
(development-plan.md Phase 2, P2.8).

P1.6/ADR-0005 already computed these once, by hand: a mid-2021 parish
population figure (baked into a reference JSON) and a Census Day OA-level
population/household sum for Cholsey/Moulsford (also baked into a
reference JSON, pulled from nomis once and never re-fetched). This module
turns both into genuinely re-fetchable live sources, per ADR-0005's own
note that this remained P2.8's job.

Two independent sources:

1. **Direct parish-level mid-year population** (ONS ad-hoc release,
   "Parish population estimates for mid-YYYY, based on best-fitting of
   output areas to parishes"). Covers Cholsey and all 8 comparators
   directly, no OA lookup needed. Real finding (2026-09-30): a mid-2022
   vintage now exists in addition to the mid-2021 one P1.6 used (Cholsey:
   4,423 for mid-2022 vs 4,404 for mid-2021) -- confirmed live by
   searching ons.gov.uk, not assumed. Each vintage gets a new numeric
   ad-hoc ID and isn't linked from the previous vintage's page, so this
   discovers the current one by scraping ONS's own site search for the
   newest "...formidYYYYbasedonbestfittingofoutputareastoparishes" result,
   the same "search, don't guess" pattern as P2.5's GOV.UK site-search
   discovery for `desnz_regional_la_energy`.
2. **Census 2021 OA-level population/household counts via the nomis API**
   (NM_2021_1 / NM_2059_1 -- the same two datasets P1.6's reference JSON
   was a one-off snapshot of), aggregated to parish via
   `geography.denominators.sum_by_parish` (the existing P1.4/P1.6 OA-
   >parish best-fit join, reused rather than duplicated). Scoped to
   Cholsey/Moulsford's existing OA set only, same as P1.6/ADR-0005 --
   comparator-level OA lookups remain the separate backlog item
   ("comparator-level apportionment", STATUS.md), not P2.8's to build.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from io import StringIO
from typing import Literal
from urllib.parse import urlencode

import requests
from openpyxl import load_workbook

from cholsey_pipeline.contracts import SOURCE_CONTRACTS, validate_row_count, validate_schema
from cholsey_pipeline.fetch.http import fetch_file, previous_row_count, record_row_count

ONS_SEARCH_URL = "https://www.ons.gov.uk/search"
"""Verified live 2026-09-30 (P2.8): a plain GET with a `q` query param
returns a server-rendered results page (no separate JSON API found)."""

PARISH_POP_LANDING_PATTERN = re.compile(
    r'href="(/peoplepopulationandcommunity/populationandmigration/populationestimates/'
    r"adhocs/\d+parishpopulationestimatesformid(\d{4})basedonbestfittingofoutputareastoparishes)\""
)

ASSET_LINK_PATTERN = re.compile(r'href="(/file\?uri=[^"]+\.xlsx)"')

PARISH_SHEET_NAME = "Parish Populations"

_PARISH_CODE_COLUMN_PATTERN = re.compile(r"^PAR\d{2}CD$")
_PARISH_NAME_COLUMN_PATTERN = re.compile(r"^PAR\d{2}NM$")


class OnsDiscoveryError(RuntimeError):
    """Raised when an ons.gov.uk page doesn't contain what was expected --
    never guessed, so a page layout/URL change surfaces loudly."""


def find_latest_parish_population_landing_path(search_page_html: str) -> tuple[int, str]:
    """Find the newest "parish population estimates for mid-YYYY" ad-hoc
    release's landing page path in ONS site-search results HTML. Pure
    function -- tested against a committed fixture, not fetched live
    (development-plan.md §5.1)."""
    matches = PARISH_POP_LANDING_PATTERN.findall(search_page_html)
    if not matches:
        raise OnsDiscoveryError(
            "No 'parish population estimates for mid-YYYY...' ad-hoc release "
            "found in the ONS search results"
        )
    path, year = max(matches, key=lambda m: int(m[1]))
    return int(year), path


def discover_parish_population_download_url(landing_page_html: str) -> str:
    """Find the current `.xlsx` asset link on a parish-population ad-hoc
    release's landing page HTML. Pure function, same pattern as
    fetch.desnz_lsoa_energy.discover_download_url."""
    match = ASSET_LINK_PATTERN.search(landing_page_html)
    if not match:
        raise OnsDiscoveryError("No .xlsx download link found on the landing page")
    return f"https://www.ons.gov.uk{match.group(1)}"


def fetch_current_parish_population_url() -> tuple[int, str]:
    """Live: search ons.gov.uk for the current vintage, then scrape its
    landing page for the current download link."""
    search_response = requests.get(
        ONS_SEARCH_URL,
        params={"q": "parish population estimates best-fitting output areas to parishes"},
        timeout=30,
    )
    search_response.raise_for_status()
    year, landing_path = find_latest_parish_population_landing_path(search_response.text)
    landing_response = requests.get(f"https://www.ons.gov.uk{landing_path}", timeout=30)
    landing_response.raise_for_status()
    return year, discover_parish_population_download_url(landing_response.text)


@dataclass(frozen=True)
class ParishPopulationRecord:
    """One parish's total mid-year population estimate. This ad-hoc
    release is parish-level only -- district/national comparator totals
    are a separate ONS release, same "not actually the same publication"
    pattern P2.5 found for DESNZ's LSOA vs regional/LA energy data."""

    vintage_year: int
    parish_code: str
    parish_name: str
    total_population: int


def _find_column(header: tuple[object, ...], pattern: re.Pattern[str], label: str) -> int:
    for index, value in enumerate(header):
        if isinstance(value, str) and pattern.match(value):
            return index
    raise OnsDiscoveryError(
        f"No column matching {pattern.pattern!r} ({label}) found in the Parish "
        "Populations sheet header"
    )


def parse_parish_population_sheet(
    rows: list[tuple[object, ...]], vintage_year: int, parish_codes: set[str] | None = None
) -> list[ParishPopulationRecord]:
    """Parse the "Parish Populations" sheet's rows (as `(PAR*CD, PAR*NM,
    Total, ...age band columns not needed here)` tuples -- real columns
    verified live 2026-09-30), optionally filtered to `parish_codes`.

    The parish code/name columns are matched by pattern (`PAR22CD`,
    `PAR23CD`, ...) rather than a hardcoded `"PAR22CD"` -- this fetcher
    always discovers the *current* vintage (see
    `fetch_current_parish_population_url`), and ONS keys each vintage's
    columns to whichever parish-boundary edition was current when it was
    published, so a later mid-2023+ release would otherwise raise
    ValueError on a hardcoded column name that no longer exists.

    Pure function -- tested against a real trimmed fixture.
    """
    header, *data_rows = rows
    code_idx = _find_column(header, _PARISH_CODE_COLUMN_PATTERN, "parish code")
    name_idx = _find_column(header, _PARISH_NAME_COLUMN_PATTERN, "parish name")
    total_idx = header.index("Total")
    records = []
    for row in data_rows:
        code = row[code_idx]
        if code is None:
            continue  # a trailing blank/footnote row, not a real parish
        if parish_codes is not None and code not in parish_codes:
            continue
        records.append(
            ParishPopulationRecord(
                vintage_year=vintage_year,
                parish_code=code,
                parish_name=row[name_idx],
                total_population=int(row[total_idx]),
            )
        )
    return records


def read_parish_population(
    xlsx_path, vintage_year: int, parish_codes: set[str] | None = None
) -> list[ParishPopulationRecord]:
    """Live: read a downloaded parish-population workbook."""
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    worksheet = workbook[PARISH_SHEET_NAME]
    rows = [row for row in worksheet.iter_rows(values_only=True)]
    return parse_parish_population_sheet(rows, vintage_year, parish_codes)


def fetch_parish_population(
    parish_codes: set[str] | None = None,
) -> list[ParishPopulationRecord]:
    """Full live fetch: discover the current vintage's download URL, fetch
    it (via fetch.http.fetch_file), and parse rows for the requested
    parishes. Validates the result against `contracts.SOURCE_CONTRACTS`
    before returning (development-plan.md P2.10)."""
    manifest_source_id = "ons_parish_population"
    year, download_url = fetch_current_parish_population_url()
    previous_count = previous_row_count(manifest_source_id)
    result = fetch_file(
        manifest_source_id,
        download_url,
        dest_filename=f"parish_population_mid{year}.xlsx",
    )
    if result.file_path is None:
        raise RuntimeError(f"fetch_file returned no file_path for {manifest_source_id}")
    records = read_parish_population(result.file_path, year, parish_codes)
    contract = SOURCE_CONTRACTS["ons_parish_population"]
    validate_schema(records, contract)
    validate_row_count(contract, len(records), previous_count)
    record_row_count(manifest_source_id, len(records))
    return records


# --- Census 2021 OA-level counts via nomis (population, households) ---

NOMIS_URL_TEMPLATE = "https://www.nomisweb.co.uk/api/v01/dataset/{dataset_id}.data.csv"

NomisMetric = Literal["population", "households"]

NOMIS_DATASETS: dict[NomisMetric, dict[str, str]] = {
    "population": {
        "dataset_id": "NM_2021_1",
        "category_code_field": "C2021_RESTYPE_3_CODE",
        "total_category_code": "0",
        # Matches config/sources.yaml's registry id, so the manifest can be
        # joined back to that entry's licence/attribution (finding from the
        # Phase 2 PR review: earlier fetchers used ad-hoc ids that didn't).
        "registry_id": "ons_census2021_ts001_population",
    },
    "households": {
        "dataset_id": "NM_2059_1",
        "category_code_field": "C2021_HH_1_CODE",
        "total_category_code": "0",
        "registry_id": "ons_census2021_ts041_households",
    },
}
"""Real dataset IDs and category codes verified live 2026-09-30 against
the nomis API -- the same two datasets P1.6's reference JSON snapshot was
pulled from by hand (see its `_provenance` block)."""


def parse_nomis_oa_csv(csv_text: str, metric: NomisMetric) -> dict[str, int]:
    """Parse a nomis dataset CSV response into `{OA21CD: count}`, keeping
    only the "total" category row per geography (population rows also
    include "lives in a household"/"lives in a communal establishment"
    breakdown rows that must not be summed in as well).

    Pure function -- tested against a real trimmed fixture.
    """
    spec = NOMIS_DATASETS[metric]
    reader = csv.DictReader(StringIO(csv_text))
    counts: dict[str, int] = {}
    for row in reader:
        if row[spec["category_code_field"]] != spec["total_category_code"]:
            continue
        counts[row["GEOGRAPHY_CODE"]] = int(row["OBS_VALUE"])
    return counts


def fetch_oa_counts(metric: NomisMetric, oa_codes: list[str]) -> dict[str, int]:
    """Live: query the nomis API for `metric` (population or households)
    for the given OA21CD codes, via `fetch.http.fetch_file` for the same
    provenance (manifest, sha256, retrieved_at) as every other Phase 2
    source -- CLAUDE.md's "provenance on every value" rule."""
    spec = NOMIS_DATASETS[metric]
    manifest_source_id = spec["registry_id"]
    base_url = NOMIS_URL_TEMPLATE.format(dataset_id=spec["dataset_id"])
    query = urlencode({"geography": ",".join(oa_codes), "measures": "20100"})
    full_url = f"{base_url}?{query}"
    result = fetch_file(manifest_source_id, full_url, dest_filename=f"oa_{metric}.csv")
    if result.file_path is None:
        raise RuntimeError(f"fetch_file returned no file_path for {manifest_source_id}")
    return parse_nomis_oa_csv(result.file_path.read_text(encoding="utf-8"), metric)
